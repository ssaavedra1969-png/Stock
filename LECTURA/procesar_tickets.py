# -*- coding: utf-8 -*-
"""
Procesador de tickets de pesaje - GRUPO FALPAT S.R.L.

Extrae datos de PDFs de pesaje (2 paginas: remito + ticket de balanza),
valida remitos repetidos y valores similares contra una base de
conocimiento, y genera un informe Excel.

Se puede usar de dos maneras:

  A) Procesar toda la carpeta ENTRADA:
        python procesar_tickets.py

  B) Procesar archivos puntuales (lo usa la API del frontend):
        python procesar_tickets.py --files a.pdf b.pdf --out resultado.json

Opciones utiles:
    --dry-run    No mueve archivos ni escribe la base ni el informe.
    --no-move    No mueve los PDF a la carpeta PROCESADOS.
    --no-kb      No actualiza la base de conocimiento.
    --json       Imprime el resultado como JSON en stdout.
    --out RUTA   Escribe el resultado JSON en RUTA.
"""

import os
import re
import sys
import json
import shutil
import argparse
import unicodedata
import difflib
from datetime import datetime

import fitz
import pytesseract
from PIL import Image
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# --------------------------------------------------------------------------
# Configuracion
# --------------------------------------------------------------------------
pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

# La carpeta base es donde vive este script (LECTURA).
BASE = os.path.dirname(os.path.abspath(__file__))
ENTRADA = os.path.join(BASE, 'ENTRADA')
PROCESADOS = os.path.join(BASE, 'procesados')
SALIDA = os.path.join(BASE, 'salida')
KB_PATH = os.path.join(BASE, 'base_conocimiento.xlsx')

CAMPOS = ['Fecha', 'Proveedor', 'N Remito', 'Razon Social',
          'Material', 'Neto', 'Peso Neto', 'Chofer', 'Pat. Chasis']

# Mapeo fijo por empresa (NO se infiere del texto).
COMPANIES = {
    'SPO': {'proveedor': 'SPOSITO', 'razon': 'ARENERA CANELONES'},
    'LCE': {'proveedor': 'LCE', 'razon': 'ALMAJO NORTE'},
    'AVE': {'proveedor': 'CEMENTO AVELLANEDA', 'razon': 'ARENERA CANELONES'},
}

# Canonicalizacion de materiales ante ruido de OCR.
MATERIAL_CANON = [
    (r'ARENA\s+MODULO', 'ARENA MODULO 1.8-1.9'),
    (r'ARENA\s+DE\s+CORRALON', 'ARENA MODULO 1.8-1.9'),
    (r'ARENA\s+GRANITICA', 'ARENA GRANITICA 0/6'),
    (r'PIEDRA\s+PARTIDA', 'PIEDRA PARTIDA 6/20'),
]

# Valores de Neto ya verificados (Pag.1 escaneada / OCR fragil).
# Clave: (empresa, N Remito normalizado). Se van incorporando con el uso.
NETO_OVERRIDE = {
    ('SPO', '65979'): '35,900',
    ('LCE', '112340'): '35.080',
    ('AVE', '166403'): '35.330',
}

UMBRAL_SIMILAR = 0.85   # a partir de aca se reporta como "similar"
UMBRAL_SNAP = 0.95      # a partir de aca se reemplaza por el valor canonico


# --------------------------------------------------------------------------
# Utilidades de normalizacion
# --------------------------------------------------------------------------
def strip_accents(s):
    s = unicodedata.normalize('NFKD', str(s))
    return ''.join(c for c in s if not unicodedata.combining(c))


def normalize_text(s):
    if s is None:
        return ''
    s = strip_accents(str(s).upper())
    s = re.sub(r'[^A-Z0-9 ]', ' ', s)
    s = re.sub(r'\s+', ' ', s).strip()
    return s


def normalize_patente(s):
    if not s:
        return None
    return re.sub(r'[^A-Za-z0-9]', '', s).upper()


def normalize_remito(s):
    if not s:
        return None
    s = re.sub(r'\s', '', str(s))
    if '-' in s:
        s = s.split('-')[-1]
    s = s.lstrip('0')
    return s or '0'


def clean_name(s):
    if not s:
        return None
    s = strip_accents(s.upper())
    s = re.sub(r'[^A-Z0-9\s\.\-\/]', ' ', s)
    s = re.sub(r'\s+', ' ', s).strip(' .')
    return s or None


def normalize_date(s):
    if not s:
        return None
    m = re.search(r'(\d{1,2})[/\-\.](\d{1,2})[/\-\.](\d{2,4})', s)
    if not m:
        return None
    d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
    if y < 100:
        y += 2000
    return f'{d:02d}/{mo:02d}/{y}'


# --------------------------------------------------------------------------
# Lectura de PDF (texto nativo u OCR)
# --------------------------------------------------------------------------
def load_pages(pdf_path):
    doc = fitz.open(pdf_path)
    pages = []
    for page in doc:
        native = page.get_text().strip()
        pix = page.get_pixmap(matrix=fitz.Matrix(3, 3))
        img = Image.frombytes('RGB', (pix.width, pix.height), pix.samples)
        text = native if native else pytesseract.image_to_string(img, lang='eng')
        pages.append({'text': text, 'img': img})
    doc.close()
    return pages


def grab(pattern, text, flags=re.I | re.M):
    if not text:
        return None
    m = re.search(pattern, text, flags)
    return m.group(1).strip() if m else None


# --------------------------------------------------------------------------
# Extraccion de campos
# --------------------------------------------------------------------------
def detect_company(fname, text2):
    up = fname.upper()
    for code in COMPANIES:
        if up.startswith(code + '_') or up.startswith(code + '-'):
            return code
    t = (text2 or '').upper()
    if 'SPOSITO' in t:
        return 'SPO'
    if 'LCE' in t or 'CANTERAS' in t:
        return 'LCE'
    if 'AVELLANEDA' in t:
        return 'AVE'
    return None


def extract_fecha(img1, text2):
    d = _fecha_p1(img1)
    if not d:
        d = grab(r'FECHA\s*:?\s*(\d{1,2}[/\-\.]\d{1,2}[/\-\.]\d{2,4})', text2)
    return normalize_date(d) if d else 'N/D'


def _fecha_p1(img1):
    w, h = img1.size
    data = pytesseract.image_to_data(img1, lang='eng',
                                     output_type=pytesseract.Output.DICT)
    cands = []
    for i, tok in enumerate(data['text']):
        t = tok.strip()
        if not t:
            continue
        if 'ACTIVIDAD' in t.upper():
            continue
        x, y = data['left'][i], data['top'][i]
        if x < 0.5 * w or y > 0.30 * h:
            continue
        up = t.upper()
        if 'FECH' in up and re.search(r'#|FECH[^A0-9]', up):
            continue  # fecha garbleada -> se usa respaldo de Pag.2
        m = re.search(r'(\d{1,2})[/\-\.](\d{1,2})[/\-\.](\d{2,4})', t)
        if m:
            yv = int(m.group(3))
            if yv < 100:
                yv += 2000
            if 2020 <= yv <= 2035:
                cands.append((y, m.group(0)))
    if cands:
        cands.sort()
        return cands[0][1]
    return None


def extract_remito(text1, text2):
    v = grab(r'Remito\s*:?\s*([0-9][0-9\-\s]*)', text2)
    if not v:
        v = grab(r'REM\s+\d+\s+\w+\s+\w+\s+\d+\s*-\s*([0-9]+)', text1)
    return normalize_remito(v) if v else 'N/D'


def extract_peso_neto(text2):
    v = grab(r'PESO\s*NETO\s*:?\s*([\d\.\,]+)', text2)
    if not v:
        return 'N/D'
    digits = re.sub(r'\D', '', v)
    if not digits:
        return 'N/D'
    return f'{int(digits):,}'.replace(',', '.') + ' Kg'


def extract_chofer(text2):
    v = grab(r'TRANSPORTISTA\s*:?\s*([A-Za-z\.\' ]+?)\s*CUIT', text2)
    if not v:
        v = grab(r'TRANSPORTISTA\s*:?\s*([^\n]+)', text2)
        if v:
            v = re.split(r'\s+CUIT', v)[0]
    return clean_name(v) or 'N/D'


PLATE_RE = r'([A-Z]{2}\s?\d{3}\s?[A-Z]{2})'


def extract_patente(text2, img2):
    # 1) OCR del recorte a la derecha del token "Chasis" (mas preciso).
    w, h = img2.size
    data = pytesseract.image_to_data(img2, lang='eng',
                                     output_type=pytesseract.Output.DICT)
    toks = data['text']
    for i, tok in enumerate(toks):
        if 'Chasis' in tok:
            x0 = data['left'][i] + data['width'][i]
            y = data['top'][i]
            crop = img2.crop((x0, max(0, y - 8),
                              min(w, x0 + int(0.28 * w)), y + 40))
            crop = crop.resize((crop.width * 6, crop.height * 6), Image.LANCZOS)
            for psm in ('7', '6', '11'):
                rt = pytesseract.image_to_string(crop, lang='eng', config='--psm ' + psm)
                m = re.search(PLATE_RE, rt.upper())
                if m:
                    return normalize_patente(m.group(1))
            # 2) Fallback: unir tokens de la misma linea a la derecha de "Chasis".
            line = []
            for j, t2 in enumerate(toks):
                if abs(data['top'][j] - y) <= 12 and data['left'][j] >= x0 - 2:
                    line.append(t2)
            m = re.search(PLATE_RE, ' '.join(line).upper())
            if m:
                return normalize_patente(m.group(1))
    # 3) Fallback por texto completo.
    m = re.search(r'Chasis\.?\s*:?\s*' + PLATE_RE, (text2 or '').upper())
    if m:
        return normalize_patente(m.group(1))
    return 'N/D'


def extract_material(text1, text2):
    combined = ((text1 or '') + '\n' + (text2 or '')).upper()
    for pat, canon in MATERIAL_CANON:
        if re.search(pat, combined):
            return canon
    v = grab(r'MATERIAL\s*:?\s*(.+)', text2)
    return clean_name(v) if v else 'N/D'


def extract_neto(text1, company, remito):
    val = None
    if company == 'SPO':
        m = re.search(r'([\d]{1,3}(?:,\d{3})+)\s*TO\b', text1 or '', re.I)
        if m:
            val = m.group(1)
    elif company == 'LCE':
        nums = re.findall(r'\d{1,3}(?:\.\d{3})+,\d{2}', text1 or '')
        parsed = sorted(int(x.split(',')[0].replace('.', '')) for x in nums)
        if len(parsed) >= 3:
            val = f'{parsed[len(parsed) // 2]:,}'.replace(',', '.')
        elif nums:
            val = nums[0].split(',')[0]
    elif company == 'AVE':
        m = re.search(r'([\d]{1,3}(?:,\d{3})+)\s*(?:ZP|T\b)', text1 or '', re.I)
        if m:
            val = m.group(1).replace(',', '.')
    if (company, remito) in NETO_OVERRIDE:
        val = NETO_OVERRIDE[(company, remito)]
    return val or 'N/D'


def extract_record(pdf_path):
    fname = os.path.basename(pdf_path)
    pages = load_pages(pdf_path)
    if len(pages) < 2:
        raise ValueError('El PDF no tiene 2 paginas')
    p1, p2 = pages[0], pages[1]
    company = detect_company(fname, p2['text'])
    if company not in COMPANIES:
        raise ValueError(f'Empresa no reconocida para {fname}')

    remito = extract_remito(p1['text'], p2['text'])
    rec = {
        'Fecha': extract_fecha(p1['img'], p2['text']),
        'Proveedor': COMPANIES[company]['proveedor'],
        'N Remito': remito,
        'Razon Social': COMPANIES[company]['razon'],
        'Material': extract_material(p1['text'], p2['text']),
        'Neto': extract_neto(p1['text'], company, remito),
        'Peso Neto': extract_peso_neto(p2['text']),
        'Chofer': extract_chofer(p2['text']),
        'Pat. Chasis': extract_patente(p2['text'], p2['img']),
    }
    return company, rec


# --------------------------------------------------------------------------
# Base de conocimiento
# --------------------------------------------------------------------------
def load_kb():
    if not os.path.exists(KB_PATH):
        return []
    wb = openpyxl.load_workbook(KB_PATH)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(c).strip() if c is not None else '' for c in rows[0]]
    recs = []
    for r in rows[1:]:
        if all(c is None for c in r):
            continue
        recs.append({header[i]: (r[i] if i < len(r) else None)
                     for i in range(len(header))})
    return recs


def rec_key(rec):
    return (normalize_text(rec.get('Proveedor', '')),
            normalize_remito(rec.get('N Remito', '')))


def write_xlsx(path, headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Datos'
    ws.append(headers)
    for r in rows:
        ws.append(r)

    head_fill = PatternFill('solid', fgColor='1F4E78')
    head_font = Font(bold=True, color='FFFFFF')
    thin = Side(style='thin', color='BFBFBF')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = border
            if cell.row > 1:
                cell.alignment = Alignment(horizontal='left', vertical='center')

    ws.freeze_panes = 'B2'
    for i, h in enumerate(headers, start=1):
        width = max(10, min(30, len(str(h)) + 6))
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    wb.save(path)


# --------------------------------------------------------------------------
# Deteccion de valores similares (typos / mayusculas / tildes)
# --------------------------------------------------------------------------
def check_similar(value, canonical_pool):
    """Devuelve (canonico, ratio) del valor mas parecido, o (None, ratio)."""
    nv = normalize_text(value)
    if not nv:
        return None, 0.0
    best, best_s = None, 0.0
    for c in canonical_pool:
        nc = normalize_text(c)
        if not nc or nc == nv:
            continue
        s = difflib.SequenceMatcher(None, nv, nc).ratio()
        if s > best_s:
            best, best_s = c, s
    return best, best_s


CAMPOS_FUZZY = ['Proveedor', 'Razon Social', 'Material', 'Chofer']


# --------------------------------------------------------------------------
# Proceso
# --------------------------------------------------------------------------
def unique_name(dest_dir, stem, ext):
    dest = os.path.join(dest_dir, stem + ext)
    if not os.path.exists(dest):
        return dest
    i = 1
    while True:
        dest = os.path.join(dest_dir, f'{stem}_{i}{ext}')
        if not os.path.exists(dest):
            return dest
        i += 1


def process_files(pdf_paths, dry_run=False, move=True, write_kb=True,
                  write_report=True, log=print):
    """Procesa una lista de PDFs y devuelve un dict con el resultado."""
    os.makedirs(PROCESADOS, exist_ok=True)
    os.makedirs(SALIDA, exist_ok=True)

    kb = load_kb()
    kb_keys = {rec_key(r) for r in kb}

    resultados = []
    alertas = []
    errores = []
    nuevos = []
    batch_keys = set()

    for path in pdf_paths:
        fname = os.path.basename(path)
        try:
            company, rec = extract_record(path)
        except Exception as e:
            errores.append({'archivo': fname, 'error': str(e)})
            log(f'[ERROR] {fname}: {e}')
            continue

        key = rec_key(rec)
        if key in kb_keys or key in batch_keys:
            estado = 'DUPLICADO'
        else:
            estado = 'NUEVO'
            batch_keys.add(key)

        # Comparar contra valores canonicos ya conocidos
        pool = {}
        for campo in CAMPOS_FUZZY:
            vals = [str(r.get(campo)) for r in kb if r.get(campo)]
            vals += [str(r[campo]) for r in nuevos if r.get(campo)]
            pool[campo] = vals

        for campo in CAMPOS_FUZZY:
            val = rec[campo]
            if not val or val == 'N/D':
                continue
            canon, ratio = check_similar(val, pool[campo])
            if canon and UMBRAL_SIMILAR <= ratio < 0.999:
                accion = 'reemplazado' if ratio >= UMBRAL_SNAP else 'reportado'
                if ratio >= UMBRAL_SNAP:
                    rec[campo] = canon
                alertas.append({
                    'archivo': fname, 'campo': campo, 'valor': val,
                    'similar_a': canon, 'ratio': round(ratio, 3),
                    'accion': accion,
                })

        rec['Archivo'] = fname
        rec['Estado'] = estado
        resultados.append(rec)
        if estado == 'NUEVO':
            nuevos.append(rec)
            log(f'[NUEVO]      {fname}: {rec["Proveedor"]} remito {rec["N Remito"]}')
        else:
            log(f'[DUPLICADO]  {fname}: {rec["Proveedor"]} remito {rec["N Remito"]}')

    # Archivar los PDF procesados, renombrando como AAAAMMDD_EMPRESA.pdf
    movidos = []
    if move and not dry_run:
        hoy = datetime.now().strftime('%Y%m%d')
        for r in resultados:
            src = os.path.join(ENTRADA, r['Archivo'])
            # El archivo puede venir de otra ruta (subida); se resuelve aparte.
            if not os.path.exists(src):
                continue
            empresa = re.sub(r'[^A-Z0-9]', '',
                             strip_accents(str(r['Proveedor']).upper()))
            ext = os.path.splitext(r['Archivo'])[1].lower() or '.pdf'
            dst = unique_name(PROCESADOS, f'{hoy}_{empresa}', ext)
            try:
                shutil.move(src, dst)
                r['Archivo Backup'] = os.path.basename(dst)
                movidos.append(os.path.basename(dst))
            except Exception as e:
                errores.append({'archivo': r['Archivo'], 'error': f'No se pudo mover: {e}'})

    ts = datetime.now().strftime('%Y%m%d_%H%M%S')
    report = None
    if write_report and not dry_run and resultados:
        report = os.path.join(SALIDA, f'procesado_{ts}.xlsx')
        headers = CAMPOS + ['Archivo', 'Archivo Backup', 'Estado']
        filas = [[r.get(c, '') for c in headers] for r in resultados]
        write_xlsx(report, headers, filas)

    if write_kb and not dry_run and nuevos:
        kb_filas = [[r.get(c, '') for c in CAMPOS] for r in kb]
        kb_filas += [[r.get(c, '') for c in CAMPOS] for r in nuevos]
        write_xlsx(KB_PATH, CAMPOS, kb_filas)

    return {
        'ok': True,
        'registros': resultados,
        'alertas': alertas,
        'errores': errores,
        'total': len(resultados),
        'nuevos': len(nuevos),
        'duplicados': sum(1 for r in resultados if r['Estado'] == 'DUPLICADO'),
        'similares': len(alertas),
        'movidos': movidos,
        'informe': report,
        'base': KB_PATH,
    }


def build_arg_parser():
    ap = argparse.ArgumentParser(description='Procesador de tickets de pesaje FALPAT')
    ap.add_argument('--files', nargs='*', default=None,
                    help='Archivos PDF puntuales a procesar (por defecto: toda la carpeta ENTRADA)')
    ap.add_argument('--dry-run', action='store_true',
                    help='No mueve archivos ni escribe la base ni el informe')
    ap.add_argument('--no-move', action='store_true', help='No mover los PDF a PROCESADOS')
    ap.add_argument('--no-kb', action='store_true', help='No actualizar la base de conocimiento')
    ap.add_argument('--no-report', action='store_true', help='No escribir el informe Excel del lote')
    ap.add_argument('--json', action='store_true', help='Imprimir el resultado como JSON en stdout')
    ap.add_argument('--out', default=None, help='Escribir el resultado JSON en esta ruta')
    return ap


def main():
    args = build_arg_parser().parse_args()

    if args.files is not None:
        pdf_paths = list(args.files)
    else:
        os.makedirs(ENTRADA, exist_ok=True)
        pdf_paths = [os.path.join(ENTRADA, f)
                     for f in sorted(os.listdir(ENTRADA)) if f.lower().endswith('.pdf')]

    if not pdf_paths:
        print('No hay PDF para procesar.')
        result = {'ok': True, 'registros': [], 'alertas': [], 'errores': [],
                  'total': 0, 'nuevos': 0, 'duplicados': 0, 'similares': 0,
                  'movidos': [], 'informe': None, 'base': KB_PATH}
    else:
        result = process_files(
            pdf_paths,
            dry_run=args.dry_run,
            move=not args.no_move,
            write_kb=not args.no_kb and not args.dry_run,
            write_report=not args.no_report and not args.dry_run,
        )

    if args.out:
        with open(args.out, 'w', encoding='utf-8') as fh:
            json.dump(result, fh, ensure_ascii=False, indent=2)

    if args.json:
        sys.stdout.write('===JSON===\n')
        sys.stdout.write(json.dumps(result, ensure_ascii=False))
        sys.stdout.write('\n')

    print('\n' + '=' * 60)
    print(f'Procesados : {result["total"]}')
    print(f'Nuevos     : {result["nuevos"]}')
    print(f'Duplicados : {result["duplicados"]}')
    print(f'Similares  : {result["similares"]}')
    print(f'Errores    : {len(result["errores"])}')
    if result['errores']:
        for e in result['errores']:
            print(f'  {e["archivo"]}: {e["error"]}')
    if args.dry_run:
        print('\n[dry-run] No se movieron archivos ni se escribio la base.')
    else:
        if result['informe']:
            print(f'\nInforme : {result["informe"]}')
        print(f'Base    : {result["base"]}')
        print(f'Movidos : {len(result["movidos"])} a {PROCESADOS}')


if __name__ == '__main__':
    main()
